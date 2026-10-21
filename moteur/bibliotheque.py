"""Bibliothèque : liste de classe, import des PDF par lot, rangement.

Rangement sous `<données>/bibliotheque/` :
    <cours>/<groupe>/classe.json                    liste de classe
    <cours>/<groupe>/<lab>/<étudiant>/rapport.pdf   rapport rangé
    <cours>/<groupe>/<lab>/<étudiant>/equipe.json   coéquipiers repérés sur la page titre
    <cours>/<groupe>/<lab>/<étudiant>/statut.json   {"termine": true} quand la correction est finie
    <cours>/<groupe>/<lab>/archive.json             {"archive": true} : lab rangé dans « Terminés »
    _attente/<id>.pdf + <id>.json                   PDF à trancher par Guillaume
    _ecartes/<id>.pdf + <id>.json                   PDF écartés (récupérables, jamais supprimés)
    _remplaces/<cours>_<groupe>_<lab>_<étudiant>_<horodatage>/   ancien dossier d'un rapport remplacé, au complet
"""

import csv
import io
import re
import unicodedata
import shutil
import uuid
from datetime import datetime
from pathlib import Path

from fastapi import APIRouter, File, Form, HTTPException, UploadFile
from fastapi.responses import FileResponse
from pydantic import BaseModel, Field
from pypdf import PdfReader

from . import stockage
from .config import dossier_donnees

router = APIRouter(prefix="/api/bibliotheque")

MOTIF_LISTE = re.compile(r"ListeEtudiants_cours(?P<cours>[A-Za-z0-9]+)_gr(?P<groupe>[0-9]+)", re.I)
SEGMENT = re.compile(r"^[A-Za-z0-9_-]{1,40}$")


def racine() -> Path:
    chemin = dossier_donnees() / "bibliotheque"
    chemin.mkdir(parents=True, exist_ok=True)
    return chemin


def normaliser(texte: str) -> str:
    """Minuscules, sans accents, lettres et chiffres séparés par des espaces."""
    sans_accent = unicodedata.normalize("NFD", texte)
    sans_accent = "".join(c for c in sans_accent if unicodedata.category(c) != "Mn")
    return " ".join(re.findall(r"[a-z0-9]+", sans_accent.lower()))


def slug(texte: str) -> str:
    return normaliser(texte).replace(" ", "-")


def segment_sur(valeur: str, nom: str) -> str:
    if not SEGMENT.match(valeur):
        raise HTTPException(400, f"{nom} invalide : lettres, chiffres, - et _ seulement.")
    return valeur


# --- Liste de classe --------------------------------------------------------

def lire_cours_groupe(nom_fichier: str) -> tuple[str, str]:
    m = MOTIF_LISTE.search(nom_fichier)
    if not m:
        raise HTTPException(
            400,
            "Nom de fichier non reconnu : attendu ListeEtudiants_cours<NUMÉRO>_gr<NUMÉRO>.",
        )
    return m["cours"].upper(), m["groupe"]


ENROBAGE = re.compile(r'^="(.*)"$', re.S)
ENTETES_NOM = {"nom", "nom de famille", "nom famille", "nom de l etudiant", "nom etudiant"}
ENTETES_DA = {"da", "no de da", "no da", "numero de da", "numero da"}


def nettoyer_cellule(valeur: str) -> str:
    """Retire les espaces et l'enrobage `="..."` des exports Omnivox."""
    valeur = valeur.strip()
    m = ENROBAGE.match(valeur)
    return m[1].strip() if m else valeur


def decoder_csv(contenu: bytes) -> str:
    try:
        return contenu.decode("utf-8-sig")
    except UnicodeDecodeError:
        return contenu.decode("cp1252", errors="replace")


def lignes_du_fichier(nom: str, contenu: bytes) -> list[list[str]]:
    if nom.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook

        feuille = load_workbook(io.BytesIO(contenu), read_only=True, data_only=True).active
        lignes = [["" if c is None else str(c) for c in ligne] for ligne in feuille.iter_rows(values_only=True)]
    elif nom.lower().endswith(".csv"):
        texte = decoder_csv(contenu)
        premiere = texte.lstrip().split(chr(10), 1)[0]
        delimiteur = max(";,\t", key=premiere.count) if any(d in premiere for d in ";,\t") else ";"
        lignes = list(csv.reader(io.StringIO(texte), delimiter=delimiteur))
    else:
        raise HTTPException(400, "Format non pris en charge : fichier Excel (.xlsx) ou CSV.")
    lignes = [[nettoyer_cellule(c) for c in ligne] for ligne in lignes]
    return [ligne for ligne in lignes if any(ligne)]


def extraire_etudiants(lignes: list[list[str]]) -> list[dict]:
    """Repère la ligne d'en-tête (colonnes nom et prénom) puis lit les étudiants.

    Le « No de DA » est gardé (champ `da`) s'il est présent; colonnes et lignes vides ignorées.
    """
    for i, ligne in enumerate(lignes[:15]):
        noms = [normaliser(c) for c in ligne]
        col_prenom = next((j for j, n in enumerate(noms) if n.startswith("prenom")), None)
        col_nom = next((j for j, n in enumerate(noms) if n in ENTETES_NOM), None)
        if col_prenom is None or col_nom is None:
            continue
        col_da = next((j for j, n in enumerate(noms) if n in ENTETES_DA), None)
        etudiants = []
        for ligne in lignes[i + 1:]:
            nom = ligne[col_nom] if col_nom < len(ligne) else ""
            prenom = ligne[col_prenom] if col_prenom < len(ligne) else ""
            if not nom.strip():
                continue
            etudiant = {"nom": nom.strip(), "prenom": prenom.strip()}  # l'id vient de fusionner_classe
            da = ligne[col_da].strip() if col_da is not None and col_da < len(ligne) else ""
            if da:
                etudiant["da"] = da
            etudiants.append(etudiant)
        return etudiants
    raise HTTPException(400, "Colonnes « Nom » et « Prénom » introuvables dans les premières lignes.")


def _cle_nom(e: dict) -> str:
    return normaliser(f"{e['nom']} {e['prenom']}")


def _etiquette(e: dict) -> str:
    return f"{e['prenom']} {e['nom']}".strip()


def _nouvel_id(e: dict, pris: set[str]) -> str:
    """Id jamais utilisé dans la liste (retirés compris) : nom-prénom, puis suffixe DA, puis -2, -3…"""
    base = slug(f"{e['nom']} {e['prenom']}") or "etudiant"
    candidats = [base]
    if e.get("da"):
        candidats.append(f"{base}-{slug(e['da'])}")
    for c in candidats:
        if c not in pris:
            return c
    n = 2
    while f"{base}-{n}" in pris:
        n += 1
    return f"{base}-{n}"


def fusionner_classe(existants: list[dict], nouveaux: list[dict]) -> tuple[list[dict], dict]:
    """Fusionne la liste réimportée avec `classe.json` (audit E4). L'ordre du fichier n'a aucun effet.

    Association : 1) par DA; 2) par nom et prénom, seulement s'il n'y a ni doute ni DA contradictoire
    (même nom mais DA différent : la même personne si son ancien DA a disparu de la liste, sinon un
    homonyme, donc un nouvel étudiant). Les ids existants sont gardés; un nouvel id ne reprend jamais
    celui d'un autre. Un étudiant absent de la nouvelle liste est gardé et marqué `retire`. Un cas
    ambigu (homonymes sans DA) n'est pas modifié et sort dans `conflits`, à trancher.
    """
    recap = {"ajoutes": [], "inchanges": 0, "mis_a_jour": [], "retires": [], "reactives": [], "conflits": []}
    appariement: dict[int, dict] = {}  # position dans `nouveaux` -> étudiant existant
    pris_existants: set[int] = set()   # positions dans `existants`

    par_da: dict[str, list[int]] = {}
    for k, e in enumerate(existants):
        if e.get("da"):
            par_da.setdefault(e["da"], []).append(k)
    for j, n in enumerate(nouveaux):
        k = next((k for k in par_da.get(n.get("da", ""), []) if k not in pris_existants), None)
        if k is not None:
            appariement[j] = existants[k]
            pris_existants.add(k)

    # Par nom : seulement parmi ceux qui restent sans correspondance.
    restants_n = [j for j in range(len(nouveaux)) if j not in appariement]
    restants_e = [k for k in range(len(existants)) if k not in pris_existants]
    da_nouveaux = {n["da"] for n in nouveaux if n.get("da")}
    conflits_j: set[int] = set()
    for cle in {_cle_nom(nouveaux[j]) for j in restants_n}:
        js = [j for j in restants_n if _cle_nom(nouveaux[j]) == cle]
        ks = [k for k in restants_e if _cle_nom(existants[k]) == cle]
        if not ks:
            continue  # vrai nouvel étudiant
        if len(js) == 1 and len(ks) == 1:
            n, e = nouveaux[js[0]], existants[ks[0]]
            da_e, da_n = e.get("da"), n.get("da")
            if da_e and da_n and da_e != da_n and da_e in da_nouveaux:
                continue  # son ancien DA est porté par quelqu'un d'autre dans la liste : un homonyme
            appariement[js[0]] = e
            pris_existants.add(ks[0])
        else:  # plusieurs homonymes de part ou d'autre sans DA qui les distingue
            conflits_j.update(js)
            pris_existants.update(ks)  # ces étudiants restent tels quels, pas retirés
            recap["conflits"].append({
                "nom": _etiquette(nouveaux[js[0]]),
                "raison": "Homonymes sans DA pour les distinguer : liste inchangée pour ces étudiants. "
                          "Ajouter le DA dans le fichier puis réimporter.",
            })

    pris_ids = {e["id"] for e in existants}
    sortie = [dict(e) for e in existants]
    index = {e["id"]: s for e, s in zip(existants, sortie)}
    for j, n in enumerate(nouveaux):
        if j in appariement:
            s = index[appariement[j]["id"]]
            avant = dict(s)
            s["nom"], s["prenom"] = n["nom"], n["prenom"]
            if n.get("da"):
                s["da"] = n["da"]
            reactive = s.pop("retire", None)
            if reactive:
                recap["reactives"].append(_etiquette(s))
            if s != avant:
                recap["mis_a_jour"].append(_etiquette(s))
            elif not reactive:
                recap["inchanges"] += 1
        elif j not in conflits_j:
            ident = _nouvel_id(n, pris_ids)
            pris_ids.add(ident)
            sortie.append({"id": ident, **n})
            recap["ajoutes"].append(_etiquette(n))
    for k, e in enumerate(existants):
        if k not in pris_existants and not e.get("retire"):
            sortie[k]["retire"] = True
            recap["retires"].append(_etiquette(e))
    return sortie, recap


def dossier_groupe(cours: str, groupe: str) -> Path:
    return racine() / segment_sur(cours, "cours") / segment_sur(groupe, "groupe")


def charger_classe(cours: str, groupe: str) -> dict:
    fichier = dossier_groupe(cours, groupe) / "classe.json"
    if not fichier.is_file():
        raise HTTPException(404, f"Aucune liste de classe pour {cours} groupe {groupe}.")
    classe = _classe_lisible(fichier)
    if classe is None:
        raise HTTPException(422, f"La liste de classe de {cours} groupe {groupe} est abîmée : la réimporter.")
    return classe


@router.post("/classes")
async def importer_classe(fichier: UploadFile = File(...)) -> dict:
    cours, groupe = lire_cours_groupe(fichier.filename or "")
    etudiants = extraire_etudiants(lignes_du_fichier(fichier.filename, await fichier.read()))
    if not etudiants:
        raise HTTPException(400, "La liste ne contient aucun étudiant.")
    dossier = dossier_groupe(cours, groupe)
    dossier.mkdir(parents=True, exist_ok=True)
    with stockage.verrou(dossier):
        existante = _classe_lisible(dossier / "classe.json")
        if existante is None and (dossier / "classe.json").is_file():
            raise HTTPException(422, "La liste existante est abîmée : la réparer ou la déplacer avant de réimporter.")
        fusion, recap = fusionner_classe(existante["etudiants"] if existante else [], etudiants)
        classe = {"cours": cours, "groupe": groupe, "etudiants": fusion}
        stockage.ecrire_json(dossier / "classe.json", classe, indent=1)
    return {**classe, "recap": recap}


# --- Association d'un PDF à un étudiant -------------------------------------

MOTIF_DA = re.compile(r"(?<![0-9])([0-9]{6,9})(?![0-9])")
LONGUEUR_MIN_COMPACT = 4  # en dessous, « roy » ne doit pas se coller dans un autre mot


def extraire_da(nom_fichier: str) -> str | None:
    """N° de DA dans le nom du fichier LÉA (« nom_2626310_Rapport… »), 6 à 9 chiffres."""
    m = MOTIF_DA.search(Path(nom_fichier).stem)
    return m[1] if m else None


def candidats_par_nom(nom_fichier: str, etudiants: list[dict]) -> list[dict]:
    """Étudiants dont le nom de famille figure dans le nom du fichier.

    Si plusieurs noms correspondent, on garde les plus longs (« Tremblay-Roy »
    l'emporte sur « Roy »). Le prénom départage seulement les homonymes.
    Sans résultat, on compare sans séparateurs (« Bele-Gombos » ~ « BeleGombos »).
    """
    texte = normaliser(Path(nom_fichier).stem)
    cible = f" {texte} "
    jetons = set(texte.split())

    def compact(e: dict) -> str:
        return normaliser(e["nom"]).replace(" ", "")

    trouves = [e for e in etudiants if f" {normaliser(e['nom'])} " in cible]
    if not trouves:
        trouves = [e for e in etudiants if len(compact(e)) >= LONGUEUR_MIN_COMPACT and compact(e) in jetons]
    if not trouves:
        return []
    plus_long = max(len(compact(e)) for e in trouves)
    trouves = [e for e in trouves if len(compact(e)) == plus_long]
    if len(trouves) > 1:
        par_prenom = [e for e in trouves if e["prenom"] and f" {normaliser(e['prenom'])} " in cible]
        if len(par_prenom) == 1:
            return par_prenom
    return trouves


def associer(nom_fichier: str, etudiants: list[dict]) -> tuple[list[dict], str | None]:
    """(étudiants possibles, DA lu dans le nom). Le DA gagne; sinon le nom de famille."""
    da = extraire_da(nom_fichier)
    if da:
        par_da = [e for e in etudiants if e.get("da") == da]
        if par_da:
            return par_da, da
    return candidats_par_nom(nom_fichier, etudiants), da


def candidats(nom_fichier: str, etudiants: list[dict]) -> list[dict]:
    return associer(nom_fichier, etudiants)[0]


def etudiants_du_cours(cours: str, groupe: str | None = None) -> list[dict]:
    """Étudiants de tous les groupes du cours (ou d'un seul), chacun avec son `groupe`."""
    segment_sur(cours, "cours")
    sortie = []
    for f in sorted((racine() / cours).glob("*/classe.json")):
        if groupe and f.parent.name != groupe:
            continue
        classe = _classe_lisible(f)
        if classe:
            sortie += [{**e, "groupe": f.parent.name} for e in classe["etudiants"]]
    return sortie


def _classe_lisible(f: Path) -> dict | None:
    """`classe.json` valide, ou None (fichier abîmé : ignoré, message au journal)."""
    classe = stockage.lire_json(f)
    return classe if isinstance(classe, dict) and isinstance(classe.get("etudiants"), list) else None


def classes_du_cours(cours: str) -> list[dict]:
    """Les `classe.json` du cours (un par groupe)."""
    segment_sur(cours, "cours")
    classes = (_classe_lisible(f) for f in sorted((racine() / cours).glob("*/classe.json")))
    return [c for c in classes if c]


def a_du_texte(contenu: bytes) -> bool:
    try:
        return any((p.extract_text() or "").strip() for p in PdfReader(io.BytesIO(contenu)).pages)
    except Exception:
        return False


LIGNES_ENTETE = 15  # le bloc d'en-tête / page titre : les premières lignes non vides de la page 1
ETIQUETTES_EXCLUES = ("enseignant", "professeur", "professeure", "charge", "remis a", "destine a")


def _contient(jetons: list[str], motif: list[str]) -> bool:
    n = len(motif)
    return n > 0 and any(jetons[i:i + n] == motif for i in range(len(jetons) - n + 1))


def coequipiers(contenu: bytes, etudiants: list[dict], exclu: dict | None) -> list[dict]:
    """Étudiants du cours nommés dans l'en-tête de la première page.

    Le prénom et le nom doivent être **adjacents** sur une même ligne (« Prénom Nom », « NOM Prénom »,
    « Nom, Prénom »; casse, accents et noms composés sans importance). Un nom seul, ou un prénom et
    un nom éloignés, ne comptent pas; les lignes d'enseignant non plus. Seules les premières lignes
    de la page sont lues : un étudiant cité dans l'introduction n'est pas un coéquipier."""
    try:
        brut = PdfReader(io.BytesIO(contenu)).pages[0].extract_text() or ""
    except Exception:
        return []
    lignes = [normaliser(x) for x in brut.splitlines()]
    lignes = [x for x in lignes if x][:LIGNES_ENTETE]
    lignes = [x for x in lignes if not x.startswith(ETIQUETTES_EXCLUES)]
    jetons = [x.split() for x in lignes]
    trouves = []
    for e in etudiants:
        if exclu and (e["groupe"], e["id"]) == (exclu["groupe"], exclu["id"]):
            continue
        nom, prenom = normaliser(e["nom"]).split(), normaliser(e["prenom"]).split()
        if nom and prenom and any(_contient(j, prenom + nom) or _contient(j, nom + prenom) for j in jetons):
            trouves.append({k: e[k] for k in ("id", "groupe", "nom", "prenom")})
    return trouves


def chemin_rapport(cours: str, groupe: str, lab: str, etudiant: str) -> Path:
    return dossier_groupe(cours, groupe) / segment_sur(lab, "lab") / etudiant / "rapport.pdf"


def ranger(contenu: bytes, cours: str, groupe: str, lab: str, etudiant: str, equipe: list[dict] | None = None,
           origine: str | None = None) -> str:
    cible = chemin_rapport(cours, groupe, lab, etudiant)
    cible.parent.mkdir(parents=True, exist_ok=True)
    cible.write_bytes(contenu)
    if origine:  # nom du PDF téléchargé de LÉA, repris pour nommer la copie corrigée
        stockage.ecrire_json(cible.parent / "origine.json", {"fichier": origine})
    fichier_equipe = cible.parent / "equipe.json"
    if equipe:
        stockage.ecrire_json(fichier_equipe, equipe)
    elif fichier_equipe.exists():
        fichier_equipe.unlink()  # nouveau rapport : l'ancienne équipe ne s'applique plus
    return f"{cours}/{groupe}/{lab}/{etudiant}"


@router.post("/import")
async def importer_lot(
    cours: str = Form(...),
    lab: str = Form(...),
    groupe: str = Form(""),
    fichiers: list[UploadFile] = File(...),
) -> dict:
    """Importe un lot de PDF. Sans `groupe`, chaque PDF va au groupe de son étudiant
    (le DA est unique dans le cours). Range ceux qui correspondent à un seul étudiant,
    met les autres en attente (avec la raison), refuse les PDF sans texte."""
    lab = slug(lab)
    groupe = groupe.strip() or None
    if groupe:
        charger_classe(cours, groupe)  # 404 si la liste est absente
    etudiants = etudiants_du_cours(cours, groupe)
    if not etudiants:
        raise HTTPException(404, f"Aucune liste de classe pour {cours}.")
    resultat = {"cours": cours, "lab": lab, "ranges": [], "en_attente": [], "refuses": [], "par_groupe": {}}
    for f in fichiers:
        nom = f.filename or "sans-nom.pdf"
        contenu = await f.read()
        if not contenu.startswith(b"%PDF") or not a_du_texte(contenu):
            resultat["refuses"].append({"fichier": nom, "raison": "PDF illisible ou sans texte (scan)"})
            continue
        trouves, da = associer(nom, etudiants)
        deja = len(trouves) == 1 and chemin_rapport(cours, trouves[0]["groupe"], lab, trouves[0]["id"]).exists()
        if len(trouves) == 1 and not deja:
            e = trouves[0]
            equipe = coequipiers(contenu, etudiants, e)
            ranger(contenu, cours, e["groupe"], lab, e["id"], equipe, nom)
            resultat["ranges"].append({"fichier": nom, "etudiant": e, "groupe": e["groupe"], "equipe": equipe})
            resultat["par_groupe"][e["groupe"]] = resultat["par_groupe"].get(e["groupe"], 0) + 1
            continue
        if deja:
            raison = "deja_importe"
        elif trouves:
            raison = "ambigu"
        else:
            raison = "da_inconnu" if da else "aucune_correspondance"
        ident = uuid.uuid4().hex[:12]
        dossier = racine() / "_attente"
        dossier.mkdir(exist_ok=True)
        (dossier / f"{ident}.pdf").write_bytes(contenu)
        meta = {"id": ident, "fichier": nom, "cours": cours, "groupe": groupe, "lab": lab,
                "raison": raison, "da": da, "candidats": trouves}
        stockage.ecrire_json(dossier / f"{ident}.json", meta)
        resultat["en_attente"].append(meta)
    return resultat


@router.get("/attente")
def lister_attente() -> list[dict]:
    dossier = racine() / "_attente"
    if not dossier.is_dir():
        return []
    metas = (stockage.lire_json(f) for f in sorted(dossier.glob("*.json")))
    return [m for m in metas if isinstance(m, dict)]


def mettre_de_cote(dossier: Path, prefixe: str) -> Path:
    """Déplace tout le dossier d'un rapport (PDF, annotations, fautes, statut, note…) dans
    `_remplaces/`, hors de ce que lisent le suivi, le plagiat et l'export. Ne supprime rien."""
    base = racine() / "_remplaces"
    base.mkdir(exist_ok=True)
    horodatage = datetime.now().strftime("%Y%m%d-%H%M%S")
    cible, n = base / f"{prefixe}_{horodatage}", 1
    while cible.exists():
        n += 1
        cible = base / f"{prefixe}_{horodatage}-{n}"
    shutil.move(str(dossier), str(cible))
    return cible


class Choix(BaseModel):
    etudiant: str | None = None  # id de l'étudiant; vide = écarter le fichier
    remplacer: bool = False      # confirmation explicite : mettre de côté le rapport déjà rangé
    groupe: str | None = None    # groupe de l'étudiant (sinon celui du lot)


@router.post("/attente/{ident}")
def trancher(ident: str, choix: Choix) -> dict:
    """Guillaume choisit l'étudiant ou écarte le fichier.

    Si l'étudiant a déjà un rapport pour ce lab, `remplacer` doit être vrai (sinon 409) : l'ancien
    dossier est alors déplacé au complet dans `_remplaces/` et le nouveau PDF repart à neuf.
    Un fichier écarté est déplacé dans `_ecartes/` (récupérable). Rien n'est jamais supprimé."""
    segment_sur(ident, "id")
    dossier = racine() / "_attente"
    meta_f, pdf_f = dossier / f"{ident}.json", dossier / f"{ident}.pdf"
    if not meta_f.is_file():
        raise HTTPException(404, "Fichier en attente introuvable.")
    meta = stockage.lire_json(meta_f)
    if not isinstance(meta, dict):
        raise HTTPException(422, "Fiche du fichier en attente abîmée : réimporter le PDF.")
    contenu = pdf_f.read_bytes()
    if choix.etudiant:
        groupe = choix.groupe or meta.get("groupe")
        tous = etudiants_du_cours(meta["cours"])
        e = next((x for x in tous if x["id"] == choix.etudiant and x["groupe"] == groupe), None)
        if not e:
            raise HTTPException(400, "Étudiant absent de la liste de classe.")
        existant = chemin_rapport(meta["cours"], e["groupe"], meta["lab"], e["id"]).parent
        if existant.exists():
            if not choix.remplacer:
                raise HTTPException(409, "Cet étudiant a déjà un rapport pour ce lab : confirmer le remplacement.")
            mettre_de_cote(existant, f"{meta['cours']}_{e['groupe']}_{meta['lab']}_{e['id']}")
        ranger(contenu, meta["cours"], e["groupe"], meta["lab"], e["id"], coequipiers(contenu, tous, e), meta.get("fichier"))
        pdf_f.unlink()
    else:
        ecartes = racine() / "_ecartes"
        ecartes.mkdir(exist_ok=True)
        pdf_f.replace(ecartes / f"{ident}.pdf")
        stockage.ecrire_json(ecartes / f"{ident}.json", meta)
    meta_f.unlink()
    return {"ok": True}


# --- Consultation -----------------------------------------------------------

@router.get("")
def arbre() -> list[dict]:
    """Cours > groupes > étudiants, avec les labs dont le rapport est rangé."""
    sortie = []
    for f in sorted(racine().glob("*/*/classe.json")):
        classe = _classe_lisible(f)
        if classe is None:
            continue
        for e in classe["etudiants"]:
            rapports = sorted(f.parent.glob(f"*/{e['id']}/rapport.pdf"))
            e["labs"] = [p.parent.parent.name for p in rapports]
            equipes = {p.parent.parent.name: stockage.lire_json(p.parent / "equipe.json", [])
                       for p in rapports if (p.parent / "equipe.json").is_file()}
            if equipes:
                e["equipes"] = equipes
            e["termines"] = [p.parent.parent.name for p in rapports if _termine(p.parent)]
            notes = {p.parent.parent.name: n for p in rapports if (n := lire_note(p.parent))}
            if notes:
                e["notes"] = notes
        classe["labs_archives"] = sorted(a.parent.name for a in f.parent.glob("*/archive.json") if _lire_drapeau(a, "archive"))
        sortie.append(classe)
    return sortie


def _lire_drapeau(fichier: Path, cle: str) -> bool:
    doc = stockage.lire_json(fichier, {})
    return bool(doc.get(cle)) if isinstance(doc, dict) else False


def _termine(dossier_rapport: Path) -> bool:
    return _lire_drapeau(dossier_rapport / "statut.json", "termine")


def lire_note(dossier_rapport: Path) -> str:
    """Note saisie à « Terminé » (texte libre, ex. « 17/20 »), vide si aucune."""
    doc = stockage.lire_json(dossier_rapport / "statut.json", {})
    return str(doc.get("note") or "") if isinstance(doc, dict) else ""


class Termine(BaseModel):
    termine: bool
    note: str | None = Field(None, max_length=40)  # absent = garder la note déjà saisie


class Archive(BaseModel):
    archive: bool


class MembreEquipe(BaseModel):
    groupe: str = Field(max_length=40)
    id: str = Field(max_length=120)


class Equipe(BaseModel):
    equipe: list[MembreEquipe] = Field(max_length=20)


@router.put("/statut/{cours}/{groupe}/{lab}/{etudiant}")
def changer_statut(cours: str, groupe: str, lab: str, etudiant: str, corps: Termine) -> dict:
    """Marque un rapport terminé (ou le remet en cours). Le rapport doit exister."""
    segment_sur(etudiant, "étudiant")
    dossier = chemin_rapport(cours, groupe, lab, etudiant).parent
    if not (dossier / "rapport.pdf").is_file():
        raise HTTPException(404, "Rapport introuvable.")
    with stockage.verrou(dossier):  # la note déjà saisie est relue et réécrite sans interruption
        note = lire_note(dossier) if corps.note is None else corps.note.strip()
        stockage.ecrire_json(dossier / "statut.json", {"termine": corps.termine, "note": note})
    return {"termine": corps.termine, "note": note}


@router.put("/equipe/{cours}/{groupe}/{lab}/{etudiant}")
def changer_equipe(cours: str, groupe: str, lab: str, etudiant: str, corps: Equipe) -> dict:
    """Remplace l'équipe d'un rapport (ajout ou retrait de coéquipiers, liste vide = seul).
    Les membres doivent être des étudiants du cours; le déposant n'en fait pas partie."""
    segment_sur(etudiant, "étudiant")
    dossier = chemin_rapport(cours, groupe, lab, etudiant).parent
    if not (dossier / "rapport.pdf").is_file():
        raise HTTPException(404, "Rapport introuvable.")
    connus = {(x["groupe"], x["id"]): x for x in etudiants_du_cours(cours)}
    equipe = []
    for m in corps.equipe:
        cle = (m.groupe, m.id)
        if cle == (groupe, etudiant):
            continue
        if cle not in connus:
            raise HTTPException(400, "Coéquipier absent de la liste de classe du cours.")
        if all((x["groupe"], x["id"]) != cle for x in equipe):
            equipe.append({k: connus[cle][k] for k in ("id", "groupe", "nom", "prenom")})
    with stockage.verrou(dossier):
        if equipe:
            stockage.ecrire_json(dossier / "equipe.json", equipe)
        elif (dossier / "equipe.json").exists():
            (dossier / "equipe.json").unlink()
    return {"equipe": equipe}


@router.put("/archive/{cours}/{groupe}/{lab}")
def changer_archive(cours: str, groupe: str, lab: str, corps: Archive) -> dict:
    """Archive un lab (il passe dans « Terminés » à l'écran; rien n'est déplacé ni supprimé)."""
    dossier = dossier_groupe(cours, groupe) / segment_sur(lab, "lab")
    if not dossier.is_dir():
        raise HTTPException(404, "Lab introuvable.")
    stockage.ecrire_json(dossier / "archive.json", {"archive": corps.archive})
    return {"archive": corps.archive}


@router.get("/rapport/{cours}/{groupe}/{lab}/{etudiant}")
def rapport(cours: str, groupe: str, lab: str, etudiant: str) -> FileResponse:
    chemin = chemin_rapport(cours, groupe, lab, segment_sur(etudiant, "étudiant"))
    if not chemin.is_file():
        raise HTTPException(404, "Rapport introuvable.")
    return FileResponse(chemin, media_type="application/pdf")

